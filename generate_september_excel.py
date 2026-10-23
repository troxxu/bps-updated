import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_pa_september_excel():
    with open('staffList.json', 'r', encoding='utf-8') as f:
        staff_list = json.load(f)

    with open('pa_september_rekap.json', 'r', encoding='utf-8') as f:
        rekap = json.load(f)

    staff_map = {s['id']: s for s in staff_list}

    def get_predikat(num_avg):
        try:
            score = float(num_avg)
        except:
            return "Belum Ada Penilaian"
        if score >= 9.0:
            return "EXCELLENT (A)"
        elif score >= 7.5:
            return "Sangat Baik (AB)"
        elif score >= 6.6:
            return "Baik (B)"
        elif score >= 5.6:
            return "Cukup (C)"
        elif score >= 3.1:
            return "Kurang Mampu (D)"
        elif score > 0:
            return "BURUK (E)"
        else:
            return "Belum Ada Penilaian"

    def get_ind_avg(sc, key, fallback_key=None):
        val = sc.get(key)
        if val is None and fallback_key:
            val = sc.get(fallback_key)
        if isinstance(val, dict):
            return val.get('average', 0)
        elif isinstance(val, (int, float)):
            return float(val)
        return 0.0

    enriched = []
    for sid, sdata in rekap.items():
        meta = staff_map.get(sid, {})
        avg = sdata.get('averageScore', 0)
        scores = sdata.get('scores', {})
        
        raw_scores_list = [v.get('rawScore', v.get('average', 0)) for v in scores.values() if isinstance(v, dict)]
        raw_mean = sum(raw_scores_list) / len(raw_scores_list) if raw_scores_list else avg

        enriched.append({
            'id': sid,
            'nama': meta.get('nama', sid),
            'role': meta.get('role', ''),
            'roleCategory': meta.get('roleCategory', ''),
            'dept': meta.get('departemenBEM', ''),
            'deptFull': meta.get('departemenBEMFull', ''),
            'nrp': meta.get('nrp', ''),
            'asalDept': meta.get('asalDepartemen', ''),
            'avgScore': avg,
            'rawMean': raw_mean,
            'predikat': get_predikat(avg),
            'evaluatorCount': sdata.get('evaluatorCount', 0),
            'kabinetEvaluatorCount': sdata.get('kabinetEvaluatorCount', 0),
            'stafEvaluatorCount': sdata.get('stafEvaluatorCount', 0),
            'hasSelfAssessment': sdata.get('hasSelfAssessment', False),
            'scores': scores
        })

    staf_indicator_labels = {
        'fundamental': 'Fundamental',
        'pemahaman_proker': 'Pemahaman Proker',
        'disiplin': 'Kedisiplinan',
        'tanggung_jawab': 'Tanggung Jawab',
        'inisiatif': 'Inisiatif Kerja',
        'inovatif': 'Kreativitas & Inovasi',
        'kritis': 'Berpikir Kritis',
        'responsif': 'Responsivitas',
        'etika': 'Etika & Sikap',
        'solidaritas': 'Solidaritas Tim',
        'rasa_kekeluargaan': 'Rasa Kekeluargaan',
        'sosial': 'Kepekaan Sosial'
    }

    # Department Rankings September
    dept_scores = {}
    for e in enriched:
        d = e['dept']
        if d == 'BPH': continue
        if d not in dept_scores:
            dept_scores[d] = []
        dept_scores[d].append(e['avgScore'])

    dept_rankings = []
    for d, sc_list in dept_scores.items():
        avg_d = sum(sc_list) / len(sc_list)
        dept_rankings.append({
            'dept': d,
            'code': d,
            'score': round(avg_d, 2),
            'members': len(sc_list),
            'predikat': get_predikat(avg_d),
            'winner': False
        })
    dept_rankings.sort(key=lambda x: x['score'], reverse=True)
    for i, d in enumerate(dept_rankings, 1):
        d['rank'] = i
        if i == 1:
            d['winner'] = True
            d['desc'] = 'Peringkat 1 Se-FTEIC (Juara Best Department of the Month - PA September 2026)'
        else:
            d['desc'] = f'Peringkat {i} Se-FTEIC'

    # Staf only list sorted by score descending
    staf_list = [e for e in enriched if e['roleCategory'] in ('Staff', 'Staf')]
    staf_list.sort(key=lambda x: (round(x['avgScore'], 4), round(x['rawMean'], 6)), reverse=True)

    # Sekdep only list sorted by score descending
    sekdep_list = [e for e in enriched if e['roleCategory'] == 'Sekretaris' or 'Sekretaris Departemen' in e['role'] or 'Sekdep' in e['role']]
    sekdep_list.sort(key=lambda x: (round(x['avgScore'], 4), round(x['rawMean'], 6)), reverse=True)

    # Best staff per dept
    best_staff_dept = []
    dept_order = ['DAGRI', 'LUGRI', 'KWU', 'SOSMAS', 'PSDM', 'RISTEK']
    for d in dept_order:
        dstaff = [s for s in staf_list if s['dept'] == d]
        if dstaff:
            dstaff.sort(key=lambda x: (round(x['avgScore'], 4), round(x['rawMean'], 6)), reverse=True)
            best_staff_dept.append(dstaff[0])

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Styles
    navy_fill = PatternFill(start_color="1D536B", end_color="1D536B", fill_type="solid")
    teal_fill = PatternFill(start_color="2A7596", end_color="2A7596", fill_type="solid")
    gold_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    silver_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bronze_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    font_title = Font(name="Segoe UI", size=15, bold=True, color="1D536B")
    font_subtitle = Font(name="Segoe UI", size=10, bold=True, color="475569")
    font_section = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
    font_regular = Font(name="Segoe UI", size=9, color="0F172A")
    font_gold_bold = Font(name="Segoe UI", size=9, bold=True, color="B45309")

    thin_border_side = Side(style='thin', color='CBD5E1')
    thick_bottom_side = Side(style='medium', color='1D536B')

    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_header = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # =========================================================================
    # SHEET 1: RINGKASAN & PAPAN APRESIASI SEPTEMBER
    # =========================================================================
    ws1 = wb.create_sheet(title="Papan Apresiasi & Juara")
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:N1")
    ws1["A1"] = "REKAPITULASI RESMI & PAPAN APRESIASI PERFORMANCE APPRAISAL (PA) SEPTEMBER 2026"
    ws1["A1"].font = font_title
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws1.merge_cells("A2:N2")
    ws1["A2"] = "BEM FAKULTAS TEKNOLOGI ELEKTRO DAN INFORMATIKA CERDAS (FTEIC) ITS • KABINET POLARIS 2026"
    ws1["A2"].font = font_subtitle
    ws1["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws1.merge_cells("A3:N3")
    ws1["A3"] = "Periode Pengisian Resmi Ditutup (21 Okt 2026 23.59 WIB) • Hasil Rekapitulasi 97 Fungsionaris Lengkap"
    ws1["A3"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    ws1["A3"].alignment = Alignment(horizontal="left", vertical="center")

    curr_row = 5

    # TABEL 1: BEST DEPARTMENT
    ws1.merge_cells(f"A{curr_row}:H{curr_row}")
    ws1[f"A{curr_row}"] = "🏆 1. BEST DEPARTMENT OF THE MONTH (PA SEPTEMBER 2026)"
    ws1[f"A{curr_row}"].font = font_section
    ws1[f"A{curr_row}"].fill = navy_fill
    ws1[f"A{curr_row}"].alignment = align_left
    ws1.row_dimensions[curr_row].height = 25
    curr_row += 1

    dept_headers = ["Peringkat", "Departemen", "Kode", "Nilai Rata-rata", "Predikat Kinerja", "Jumlah Fungsionaris", "Status Partisipasi", "Keterangan & Prestasi"]
    for col_idx, h in enumerate(dept_headers, 1):
        cell = ws1.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = teal_fill
        cell.alignment = align_header
        cell.border = border_header
    ws1.row_dimensions[curr_row].height = 22
    curr_row += 1

    for d in dept_rankings:
        row_cells = [
            ws1.cell(row=curr_row, column=1, value=f"#{d['rank']}" if not d['winner'] else f"🏆 #{d['rank']}"),
            ws1.cell(row=curr_row, column=2, value=d['dept']),
            ws1.cell(row=curr_row, column=3, value=d['code']),
            ws1.cell(row=curr_row, column=4, value=d['score']),
            ws1.cell(row=curr_row, column=5, value=d['predikat']),
            ws1.cell(row=curr_row, column=6, value=f"{d['members']} Fungsionaris"),
            ws1.cell(row=curr_row, column=7, value="100% Selesai Direkap"),
            ws1.cell(row=curr_row, column=8, value=d['desc']),
        ]
        is_win = d['winner']
        for i, c in enumerate(row_cells):
            c.border = border_cell
            if is_win:
                c.fill = gold_fill
                c.font = font_gold_bold if i in (0, 1, 3) else font_bold
            else:
                c.font = font_regular if i not in (0, 3) else font_bold
                if curr_row % 2 == 0:
                    c.fill = zebra_fill
        row_cells[0].alignment = align_center
        row_cells[2].alignment = align_center
        row_cells[3].alignment = align_center
        row_cells[3].number_format = '0.00'
        row_cells[4].alignment = align_center
        row_cells[5].alignment = align_center
        row_cells[6].alignment = align_center
        ws1.row_dimensions[curr_row].height = 20
        curr_row += 1

    curr_row += 2

    # TABEL 2: BEST TOP 3 STAFF
    ws1.merge_cells(f"A{curr_row}:O{curr_row}")
    ws1[f"A{curr_row}"] = "🌟 2. BEST TOP 3 STAFF OF THE MONTH (SE-BEM FTEIC)"
    ws1[f"A{curr_row}"].font = font_section
    ws1[f"A{curr_row}"].fill = navy_fill
    ws1[f"A{curr_row}"].alignment = align_left
    ws1.row_dimensions[curr_row].height = 25
    curr_row += 1

    staff_top3_headers = [
        "Peringkat", "ID Staf", "Nama Lengkap", "Departemen", "NRP", "Asal Jurusan", 
        "Nilai Rata-rata", "Nilai Desimal", "Predikat", "Penilai", "Etika", "Kritis", "Tanggung Jawab", 
        "Pemahaman Proker", "Disiplin"
    ]
    for col_idx, h in enumerate(staff_top3_headers, 1):
        cell = ws1.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = teal_fill
        cell.alignment = align_header
        cell.border = border_header
    ws1.row_dimensions[curr_row].height = 24
    curr_row += 1

    medals = ["🥇 Juara 1", "🥈 Juara 2", "🥉 Juara 3"]
    medal_fills = [gold_fill, silver_fill, bronze_fill]

    for rank_idx, s in enumerate(staf_list[:3]):
        sc = s['scores']
        row_cells = [
            ws1.cell(row=curr_row, column=1, value=medals[rank_idx]),
            ws1.cell(row=curr_row, column=2, value=s['id']),
            ws1.cell(row=curr_row, column=3, value=s['nama']),
            ws1.cell(row=curr_row, column=4, value=s['dept']),
            ws1.cell(row=curr_row, column=5, value=s['nrp']),
            ws1.cell(row=curr_row, column=6, value=s['asalDept']),
            ws1.cell(row=curr_row, column=7, value=s['avgScore']),
            ws1.cell(row=curr_row, column=8, value=s['rawMean']),
            ws1.cell(row=curr_row, column=9, value=s['predikat']),
            ws1.cell(row=curr_row, column=10, value=f"{s['evaluatorCount']} Orang"),
            ws1.cell(row=curr_row, column=11, value=get_ind_avg(sc, 'etika')),
            ws1.cell(row=curr_row, column=12, value=get_ind_avg(sc, 'kritis')),
            ws1.cell(row=curr_row, column=13, value=get_ind_avg(sc, 'tanggung_jawab')),
            ws1.cell(row=curr_row, column=14, value=get_ind_avg(sc, 'pemahaman_proker')),
            ws1.cell(row=curr_row, column=15, value=get_ind_avg(sc, 'disiplin')),
        ]
        for idx, c in enumerate(row_cells):
            c.border = border_cell
            c.fill = medal_fills[rank_idx]
            if idx in (0, 2, 6):
                c.font = font_bold
            else:
                c.font = font_regular
        
        row_cells[0].alignment = align_center
        row_cells[1].alignment = align_center
        row_cells[3].alignment = align_center
        row_cells[4].alignment = align_center
        row_cells[6].alignment = align_center
        row_cells[6].number_format = '0.00'
        row_cells[7].alignment = align_center
        row_cells[7].number_format = '0.0000'
        row_cells[8].alignment = align_center
        row_cells[9].alignment = align_center
        for col_k in range(10, 15):
            row_cells[col_k].alignment = align_center
            row_cells[col_k].number_format = '0.00'

        ws1.row_dimensions[curr_row].height = 22
        curr_row += 1

    curr_row += 2

    # TABEL 3: BEST TOP 3 SEKDEP
    ws1.merge_cells(f"A{curr_row}:O{curr_row}")
    ws1[f"A{curr_row}"] = "📝 3. BEST TOP 3 SEKRETARIS DEPARTEMEN (SEKDEP) OF THE MONTH"
    ws1[f"A{curr_row}"].font = font_section
    ws1[f"A{curr_row}"].fill = navy_fill
    ws1[f"A{curr_row}"].alignment = align_left
    ws1.row_dimensions[curr_row].height = 25
    curr_row += 1

    sekdep_top3_headers = [
        "Peringkat", "ID Staf", "Nama Lengkap", "Departemen", "NRP", "Asal Jurusan", 
        "Nilai Rata-rata", "Nilai Desimal", "Predikat", "Penilai", "Etika", "Tanggung Jawab", "Sosial", 
        "Solidaritas", "Kedisiplinan"
    ]
    for col_idx, h in enumerate(sekdep_top3_headers, 1):
        cell = ws1.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = teal_fill
        cell.alignment = align_header
        cell.border = border_header
    ws1.row_dimensions[curr_row].height = 24
    curr_row += 1

    for rank_idx, s in enumerate(sekdep_list[:3]):
        sc = s['scores']
        row_cells = [
            ws1.cell(row=curr_row, column=1, value=medals[rank_idx]),
            ws1.cell(row=curr_row, column=2, value=s['id']),
            ws1.cell(row=curr_row, column=3, value=s['nama']),
            ws1.cell(row=curr_row, column=4, value=s['dept']),
            ws1.cell(row=curr_row, column=5, value=s['nrp']),
            ws1.cell(row=curr_row, column=6, value=s['asalDept']),
            ws1.cell(row=curr_row, column=7, value=s['avgScore']),
            ws1.cell(row=curr_row, column=8, value=s['rawMean']),
            ws1.cell(row=curr_row, column=9, value=s['predikat']),
            ws1.cell(row=curr_row, column=10, value=f"{s['evaluatorCount']} Orang"),
            ws1.cell(row=curr_row, column=11, value=get_ind_avg(sc, 'etika')),
            ws1.cell(row=curr_row, column=12, value=get_ind_avg(sc, 'tanggung_jawab')),
            ws1.cell(row=curr_row, column=13, value=get_ind_avg(sc, 'sosial')),
            ws1.cell(row=curr_row, column=14, value=get_ind_avg(sc, 'solidaritas')),
            ws1.cell(row=curr_row, column=15, value=get_ind_avg(sc, 'disiplin')),
        ]
        for idx, c in enumerate(row_cells):
            c.border = border_cell
            c.fill = medal_fills[rank_idx]
            if idx in (0, 2, 6):
                c.font = font_bold
            else:
                c.font = font_regular
        
        row_cells[0].alignment = align_center
        row_cells[1].alignment = align_center
        row_cells[3].alignment = align_center
        row_cells[4].alignment = align_center
        row_cells[6].alignment = align_center
        row_cells[6].number_format = '0.00'
        row_cells[7].alignment = align_center
        row_cells[7].number_format = '0.0000'
        row_cells[8].alignment = align_center
        row_cells[9].alignment = align_center
        for col_k in range(10, 15):
            row_cells[col_k].alignment = align_center
            row_cells[col_k].number_format = '0.00'

        ws1.row_dimensions[curr_row].height = 22
        curr_row += 1

    curr_row += 2

    # TABEL 4: BEST STAFF TIAP DEPARTEMEN
    ws1.merge_cells(f"A{curr_row}:K{curr_row}")
    ws1[f"A{curr_row}"] = "🏅 4. BEST STAFF SETIAP DEPARTEMEN (6 DEPARTEMEN RESMI BEM FTEIC)"
    ws1[f"A{curr_row}"].font = font_section
    ws1[f"A{curr_row}"].fill = navy_fill
    ws1[f"A{curr_row}"].alignment = align_left
    ws1.row_dimensions[curr_row].height = 25
    curr_row += 1

    best_dept_headers = [
        "Departemen", "ID Staf", "Nama Staf Terbaik", "NRP", "Asal Jurusan", 
        "Nilai Rata-rata", "Nilai Desimal", "Predikat", "Jumlah Penilai", "Indikator Capaian Tertinggi", "Skor Tertinggi"
    ]
    for col_idx, h in enumerate(best_dept_headers, 1):
        cell = ws1.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = teal_fill
        cell.alignment = align_header
        cell.border = border_header
    ws1.row_dimensions[curr_row].height = 24
    curr_row += 1

    for s in best_staff_dept:
        sc = s['scores']
        best_ind_name = "-"
        best_ind_val = 0
        for k, v in sc.items():
            if isinstance(v, dict):
                val = v.get('average', 0)
                if val > best_ind_val:
                    best_ind_val = val
                    best_ind_name = staf_indicator_labels.get(k, k)

        row_cells = [
            ws1.cell(row=curr_row, column=1, value=s['dept']),
            ws1.cell(row=curr_row, column=2, value=s['id']),
            ws1.cell(row=curr_row, column=3, value=s['nama']),
            ws1.cell(row=curr_row, column=4, value=s['nrp']),
            ws1.cell(row=curr_row, column=5, value=s['asalDept']),
            ws1.cell(row=curr_row, column=6, value=s['avgScore']),
            ws1.cell(row=curr_row, column=7, value=s['rawMean']),
            ws1.cell(row=curr_row, column=8, value=s['predikat']),
            ws1.cell(row=curr_row, column=9, value=f"{s['evaluatorCount']} Penilai"),
            ws1.cell(row=curr_row, column=10, value=best_ind_name),
            ws1.cell(row=curr_row, column=11, value=best_ind_val),
        ]
        for idx, c in enumerate(row_cells):
            c.border = border_cell
            if curr_row % 2 == 0:
                c.fill = zebra_fill
            if idx in (0, 2, 5):
                c.font = font_bold
            else:
                c.font = font_regular
        
        row_cells[0].alignment = align_center
        row_cells[1].alignment = align_center
        row_cells[3].alignment = align_center
        row_cells[5].alignment = align_center
        row_cells[5].number_format = '0.00'
        row_cells[6].alignment = align_center
        row_cells[6].number_format = '0.0000'
        row_cells[7].alignment = align_center
        row_cells[8].alignment = align_center
        row_cells[9].alignment = align_center
        row_cells[10].alignment = align_center
        row_cells[10].number_format = '0.00'

        ws1.row_dimensions[curr_row].height = 20
        curr_row += 1


    # =========================================================================
    # SHEET 2: DATA NILAI STAF & SEKDEP SEPTEMBER (12 INDIKATOR)
    # =========================================================================
    ws2 = wb.create_sheet(title="Data Nilai Staf & Sekdep")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:AA1")
    ws2["A1"] = "DATA LENGKAP NILAI PERFORMANCE APPRAISAL SEPTEMBER 2026 — STAF & SEKRETARIS DEPARTEMEN"
    ws2["A1"].font = font_title
    ws2["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws2.merge_cells("A2:AA2")
    ws2["A2"] = "Rubrik 12 Indikator Standar Kompetensi Staf • Skor Akhir Terbobot 360° (Kabinet 55%, Rekan 35%, Diri 10%)"
    ws2["A2"].font = font_subtitle
    ws2["A2"].alignment = Alignment(horizontal="left", vertical="center")

    headers_ws2 = [
        "No", "ID Staf", "Nama Lengkap", "NRP", "Departemen", "Program Studi", "Jabatan", "Kategori",
        "Total Penilai", "Penilai Kabinet", "Penilai Staf", "Evaluasi Diri",
        "Rata-rata Akhir", "Nilai Desimal (Raw)", "Predikat Kinerja",
        "Fundamental", "Pemahaman Proker", "Kedisiplinan", "Tanggung Jawab", "Inisiatif Kerja",
        "Kreativitas & Inovasi", "Berpikir Kritis", "Responsivitas", "Etika & Sikap", "Solidaritas Tim",
        "Rasa Kekeluargaan", "Kepekaan Sosial"
    ]
    
    row_h2 = 4
    for col_idx, h in enumerate(headers_ws2, 1):
        cell = ws2.cell(row=row_h2, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = navy_fill
        cell.alignment = align_header
        cell.border = border_header
    ws2.row_dimensions[row_h2].height = 28

    staf_and_sekdep = [e for e in enriched if e['roleCategory'] in ('Staff', 'Staf', 'Sekretaris')]
    staf_and_sekdep.sort(key=lambda x: (x['dept'], -x['avgScore'], -x['rawMean']))

    row_data = row_h2 + 1
    for no, s in enumerate(staf_and_sekdep, 1):
        sc = s['scores']
        r_cells = [
            ws2.cell(row=row_data, column=1, value=no),
            ws2.cell(row=row_data, column=2, value=s['id']),
            ws2.cell(row=row_data, column=3, value=s['nama']),
            ws2.cell(row=row_data, column=4, value=s['nrp']),
            ws2.cell(row=row_data, column=5, value=s['dept']),
            ws2.cell(row=row_data, column=6, value=s['asalDept']),
            ws2.cell(row=row_data, column=7, value=s['role']),
            ws2.cell(row=row_data, column=8, value=s['roleCategory']),
            ws2.cell(row=row_data, column=9, value=s['evaluatorCount']),
            ws2.cell(row=row_data, column=10, value=s['kabinetEvaluatorCount']),
            ws2.cell(row=row_data, column=11, value=s['stafEvaluatorCount']),
            ws2.cell(row=row_data, column=12, value="Ya" if s['hasSelfAssessment'] else "Tidak"),
            ws2.cell(row=row_data, column=13, value=s['avgScore']),
            ws2.cell(row=row_data, column=14, value=s['rawMean']),
            ws2.cell(row=row_data, column=15, value=s['predikat']),
            ws2.cell(row=row_data, column=16, value=get_ind_avg(sc, 'fundamental')),
            ws2.cell(row=row_data, column=17, value=get_ind_avg(sc, 'pemahaman_proker')),
            ws2.cell(row=row_data, column=18, value=get_ind_avg(sc, 'disiplin')),
            ws2.cell(row=row_data, column=19, value=get_ind_avg(sc, 'tanggung_jawab')),
            ws2.cell(row=row_data, column=20, value=get_ind_avg(sc, 'inisiatif')),
            ws2.cell(row=row_data, column=21, value=get_ind_avg(sc, 'inovatif')),
            ws2.cell(row=row_data, column=22, value=get_ind_avg(sc, 'kritis')),
            ws2.cell(row=row_data, column=23, value=get_ind_avg(sc, 'responsif')),
            ws2.cell(row=row_data, column=24, value=get_ind_avg(sc, 'etika')),
            ws2.cell(row=row_data, column=25, value=get_ind_avg(sc, 'solidaritas')),
            ws2.cell(row=row_data, column=26, value=get_ind_avg(sc, 'rasa_kekeluargaan')),
            ws2.cell(row=row_data, column=27, value=get_ind_avg(sc, 'sosial')),
        ]
        for idx, c in enumerate(r_cells):
            c.border = border_cell
            if row_data % 2 == 0:
                c.fill = zebra_fill
            if idx in (0, 2, 12):
                c.font = font_bold
            else:
                c.font = font_regular
        
        r_cells[0].alignment = align_center
        r_cells[1].alignment = align_center
        r_cells[3].alignment = align_center
        r_cells[4].alignment = align_center
        r_cells[7].alignment = align_center
        r_cells[8].alignment = align_center
        r_cells[9].alignment = align_center
        r_cells[10].alignment = align_center
        r_cells[11].alignment = align_center
        r_cells[12].alignment = align_center
        r_cells[12].number_format = '0.00'
        r_cells[13].alignment = align_center
        r_cells[13].number_format = '0.0000'
        r_cells[14].alignment = align_center
        
        for k_idx in range(15, 27):
            r_cells[k_idx].alignment = align_center
            r_cells[k_idx].number_format = '0.00'

        ws2.row_dimensions[row_data].height = 19
        row_data += 1


    # =========================================================================
    # SHEET 3: DATA NILAI PIMPINAN, KABIRO & BPH SEPTEMBER
    # =========================================================================
    ws3 = wb.create_sheet(title="Data Nilai Pimpinan & BPH")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:V1")
    ws3["A1"] = "DATA LENGKAP NILAI PERFORMANCE APPRAISAL SEPTEMBER 2026 — PIMPINAN, KABIRO & BPH"
    ws3["A1"].font = font_title
    ws3["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws3.merge_cells("A2:V2")
    ws3["A2"] = "Rubrik Kompetensi Kepemimpinan & Tata Kelola Organisasi (Bobot: Kabinet 35%, Staf 55%, Diri 10%)"
    ws3["A2"].font = font_subtitle
    ws3["A2"].alignment = Alignment(horizontal="left", vertical="center")

    headers_ws3 = [
        "No", "ID Fungsionaris", "Nama Lengkap", "NRP", "Departemen / Biro", "Program Studi", "Jabatan", "Kategori",
        "Total Penilai", "Penilai Kabinet", "Penilai Staf", "Evaluasi Diri",
        "Rata-rata Akhir", "Nilai Desimal", "Predikat Kinerja",
        "Kepemimpinan (Leadership)", "Pengembangan Tim (Development)", 
        "Pemberdayaan / Tata Kelola (Empowerment / Governance)", 
        "Koordinasi / Komunikasi (Coordination / Open Comm)", 
        "Representasi & Diplomasi", "Tata Kelola Administrasi", "Keteladanan & Integritas (Exemplary)"
    ]

    row_h3 = 4
    for col_idx, h in enumerate(headers_ws3, 1):
        cell = ws3.cell(row=row_h3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = navy_fill
        cell.alignment = align_header
        cell.border = border_header
    ws3.row_dimensions[row_h3].height = 28

    leaders = [e for e in enriched if e['roleCategory'] in ('Pimpinan', 'Kabiro', 'BPH')]
    leaders.sort(key=lambda x: (x['dept'], -x['avgScore'], -x['rawMean']))

    row_data3 = row_h3 + 1
    for no, s in enumerate(leaders, 1):
        sc = s['scores']
        r_cells = [
            ws3.cell(row=row_data3, column=1, value=no),
            ws3.cell(row=row_data3, column=2, value=s['id']),
            ws3.cell(row=row_data3, column=3, value=s['nama']),
            ws3.cell(row=row_data3, column=4, value=s['nrp']),
            ws3.cell(row=row_data3, column=5, value=s['dept']),
            ws3.cell(row=row_data3, column=6, value=s['asalDept']),
            ws3.cell(row=row_data3, column=7, value=s['role']),
            ws3.cell(row=row_data3, column=8, value=s['roleCategory']),
            ws3.cell(row=row_data3, column=9, value=s['evaluatorCount']),
            ws3.cell(row=row_data3, column=10, value=s['kabinetEvaluatorCount']),
            ws3.cell(row=row_data3, column=11, value=s['stafEvaluatorCount']),
            ws3.cell(row=row_data3, column=12, value="Ya" if s['hasSelfAssessment'] else "Tidak"),
            ws3.cell(row=row_data3, column=13, value=s['avgScore']),
            ws3.cell(row=row_data3, column=14, value=s['rawMean']),
            ws3.cell(row=row_data3, column=15, value=s['predikat']),
            ws3.cell(row=row_data3, column=16, value=get_ind_avg(sc, 'leadership')),
            ws3.cell(row=row_data3, column=17, value=get_ind_avg(sc, 'development')),
            ws3.cell(row=row_data3, column=18, value=get_ind_avg(sc, 'empowerment', 'governance')),
            ws3.cell(row=row_data3, column=19, value=get_ind_avg(sc, 'coordination', 'open_communication')),
            ws3.cell(row=row_data3, column=20, value=get_ind_avg(sc, 'representation')),
            ws3.cell(row=row_data3, column=21, value=get_ind_avg(sc, 'administration')),
            ws3.cell(row=row_data3, column=22, value=get_ind_avg(sc, 'exemplary')),
        ]
        for idx, c in enumerate(r_cells):
            c.border = border_cell
            if row_data3 % 2 == 0:
                c.fill = zebra_fill
            if idx in (0, 2, 12):
                c.font = font_bold
            else:
                c.font = font_regular
        
        r_cells[0].alignment = align_center
        r_cells[1].alignment = align_center
        r_cells[3].alignment = align_center
        r_cells[4].alignment = align_center
        r_cells[7].alignment = align_center
        r_cells[8].alignment = align_center
        r_cells[9].alignment = align_center
        r_cells[10].alignment = align_center
        r_cells[11].alignment = align_center
        r_cells[12].alignment = align_center
        r_cells[12].number_format = '0.00'
        r_cells[13].alignment = align_center
        r_cells[13].number_format = '0.0000'
        r_cells[14].alignment = align_center
        
        for k_idx in range(15, 22):
            r_cells[k_idx].alignment = align_center
            r_cells[k_idx].number_format = '0.00'

        ws3.row_dimensions[row_data3].height = 19
        row_data3 += 1


    # =========================================================================
    # SHEET 4: RINCIAN NILAI 360 DERAJAT (SUB-NILAI DETAIL)
    # =========================================================================
    ws4 = wb.create_sheet(title="Rincian Skor 360° (Sub-Nilai)")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:R1")
    ws4["A1"] = "RINCIAN SUB-NILAI EVALUASI 360° SETIAP INDIKATOR — PERIODE PA SEPTEMBER 2026"
    ws4["A1"].font = font_title
    ws4["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws4.merge_cells("A2:R2")
    ws4["A2"] = "Daftar Rinci Nilai Kecil: Penilaian Kabinet, Penilaian Rekan Staf, Penilaian Diri Sendiri, Jumlah Penilai & Bobot"
    ws4["A2"].font = font_subtitle
    ws4["A2"].alignment = Alignment(horizontal="left", vertical="center")

    headers_ws4 = [
        "No", "ID Fungsionaris", "Nama Lengkap", "Departemen", "Kategori Jabatan",
        "Kode Indikator", "Nama Indikator Resmi", "Nilai Akhir Indikator", "Nilai Desimal (Raw)",
        "Rata-rata Kabinet", "Jumlah Penilai Kabinet", "Bobot Kabinet (%)",
        "Rata-rata Rekan Staf", "Jumlah Penilai Rekan", "Bobot Rekan (%)",
        "Penilaian Diri (Self)", "Jumlah Self", "Bobot Self (%)"
    ]

    row_h4 = 4
    for col_idx, h in enumerate(headers_ws4, 1):
        cell = ws4.cell(row=row_h4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = navy_fill
        cell.alignment = align_header
        cell.border = border_header
    ws4.row_dimensions[row_h4].height = 28

    all_indicator_labels = {
        'leadership': 'Kepemimpinan Strategis',
        'development': 'Pengembangan Anggota',
        'empowerment': 'Pemberdayaan Tim',
        'coordination': 'Koordinasi Lembaga',
        'representation': 'Representasi & Diplomasi',
        'administration': 'Tata Kelola Administrasi',
        'exemplary': 'Keteladanan & Integritas',
        'governance': 'Tata Kelola Kabinet',
        'open_communication': 'Komunikasi Terbuka'
    }
    all_indicator_labels.update(staf_indicator_labels)

    row_data4 = row_h4 + 1
    counter_sub = 1
    for s in enriched:
        sc = s['scores']
        for ind_key, ind_val in sc.items():
            if not isinstance(ind_val, dict):
                continue
            ind_label = all_indicator_labels.get(ind_key, ind_key.replace('_', ' ').title())
            weights = ind_val.get('weights', {})

            r_cells = [
                ws4.cell(row=row_data4, column=1, value=counter_sub),
                ws4.cell(row=row_data4, column=2, value=s['id']),
                ws4.cell(row=row_data4, column=3, value=s['nama']),
                ws4.cell(row=row_data4, column=4, value=s['dept']),
                ws4.cell(row=row_data4, column=5, value=s['roleCategory']),
                ws4.cell(row=row_data4, column=6, value=ind_key),
                ws4.cell(row=row_data4, column=7, value=ind_label),
                ws4.cell(row=row_data4, column=8, value=ind_val.get('average', 0)),
                ws4.cell(row=row_data4, column=9, value=ind_val.get('rawScore', ind_val.get('average', 0))),
                ws4.cell(row=row_data4, column=10, value=ind_val.get('kabinetAvg', 0)),
                ws4.cell(row=row_data4, column=11, value=ind_val.get('kabCount', 0)),
                ws4.cell(row=row_data4, column=12, value=f"{int(weights.get('kabinet', 0)*100)}%"),
                ws4.cell(row=row_data4, column=13, value=ind_val.get('stafAvg', 0)),
                ws4.cell(row=row_data4, column=14, value=ind_val.get('stafCount', 0)),
                ws4.cell(row=row_data4, column=15, value=f"{int(weights.get('staf', 0)*100)}%"),
                ws4.cell(row=row_data4, column=16, value=ind_val.get('selfAvg', 0)),
                ws4.cell(row=row_data4, column=17, value=ind_val.get('selfCount', 0)),
                ws4.cell(row=row_data4, column=18, value=f"{int(weights.get('self', 0)*100)}%"),
            ]
            for idx, c in enumerate(r_cells):
                c.border = border_cell
                if row_data4 % 2 == 0:
                    c.fill = zebra_fill
                if idx in (0, 7):
                    c.font = font_bold
                else:
                    c.font = font_regular
            
            r_cells[0].alignment = align_center
            r_cells[1].alignment = align_center
            r_cells[3].alignment = align_center
            r_cells[4].alignment = align_center
            r_cells[5].alignment = align_center
            r_cells[7].alignment = align_center
            r_cells[7].number_format = '0.00'
            r_cells[8].alignment = align_center
            r_cells[8].number_format = '0.0000'
            r_cells[9].alignment = align_center
            r_cells[9].number_format = '0.00'
            r_cells[10].alignment = align_center
            r_cells[11].alignment = align_center
            r_cells[12].alignment = align_center
            r_cells[12].number_format = '0.00'
            r_cells[13].alignment = align_center
            r_cells[14].alignment = align_center
            r_cells[15].alignment = align_center
            r_cells[15].number_format = '0.00'
            r_cells[16].alignment = align_center
            r_cells[17].alignment = align_center

            ws4.row_dimensions[row_data4].height = 18
            row_data4 += 1
            counter_sub += 1

    # Auto-fit Column Widths for all sheets
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in (1, 2, 3):
                    continue
                val = str(cell.value or '')
                if '\n' in val:
                    val = max(val.split('\n'), key=len)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws1.column_dimensions['A'].width = 16
    ws1.column_dimensions['B'].width = 30
    ws1.column_dimensions['C'].width = 14
    ws1.column_dimensions['D'].width = 18
    ws1.column_dimensions['E'].width = 22
    ws1.column_dimensions['F'].width = 22
    ws1.column_dimensions['G'].width = 20
    ws1.column_dimensions['H'].width = 35

    ws1.freeze_panes = "A5"
    ws2.freeze_panes = "D5"
    ws3.freeze_panes = "D5"
    ws4.freeze_panes = "D5"

    output_filename = "PA_September_2026_Rekapitulasi_Lengkap.xlsx"
    wb.save(output_filename)
    print(f"Successfully generated {output_filename}!")

if __name__ == "__main__":
    build_pa_september_excel()
